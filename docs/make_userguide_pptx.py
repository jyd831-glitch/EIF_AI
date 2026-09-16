# -*- coding: utf-8 -*-
"""Generate EIF Log Tools user guide PowerPoint."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent / "EIF_Log_Tools_UserGuide.pptx"
OUT.parent.mkdir(parents=True, exist_ok=True)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(0x14, 0x32, 0x5F)
DARK = RGBColor(0x1C, 0x24, 0x30)
MUTED = RGBColor(0x5F, 0x6E, 0x7F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF0, 0xF4, 0xF8)


def set_run(run, size=18, bold=False, color=DARK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Malgun Gothic"


def add_header_bar(slide, title: str):
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, Inches(0.95))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.22), Inches(12), Inches(0.55))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = title
    set_run(run, 26, True, WHITE)


def add_footer(slide, page: int, total: int):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(7.05), Inches(12), Inches(0.3))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = f"EIF Log Tools 사용자 가이드  |  {page} / {total}"
    set_run(run, 11, False, MUTED)


def add_bullets(slide, left, top, width, height, items, size=16):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = item.get("level", 0)
        p.space_after = Pt(6)
        run = p.add_run()
        run.text = item["text"]
        set_run(run, item.get("size", size), item.get("bold", False), item.get("color", DARK))


def add_card(slide, left, top, width, height, title, lines, fill=LIGHT):
    card = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = fill
    card.line.color.rgb = RGBColor(0xC5, 0xCE, 0xD8)
    title_box = slide.shapes.add_textbox(
        Inches(left + 0.2), Inches(top + 0.15), Inches(width - 0.4), Inches(0.4)
    )
    p = title_box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = title
    set_run(run, 16, True, NAVY)
    body = slide.shapes.add_textbox(
        Inches(left + 0.2), Inches(top + 0.55), Inches(width - 0.4), Inches(height - 0.7)
    )
    tf = body.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = line
        set_run(run, 13, False, DARK)


blank = prs.slide_layouts[6]


def new_slide():
    return prs.slides.add_slide(blank)


# ---- Slide 1: Cover ----
s = new_slide()
bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = NAVY
bg.line.fill.background()
band = s.shapes.add_shape(1, 0, Inches(2.2), prs.slide_width, Inches(2.6))
band.fill.solid()
band.fill.fore_color.rgb = RGBColor(0x1A, 0x3F, 0x6E)
band.line.fill.background()
t = s.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(11.5), Inches(1))
r = t.text_frame.paragraphs[0].add_run()
r.text = "EIF Log Tools"
set_run(r, 44, True, WHITE)
t2 = s.shapes.add_textbox(Inches(0.8), Inches(3.5), Inches(11.5), Inches(0.8))
r = t2.text_frame.paragraphs[0].add_run()
r.text = "사용자 가이드 (상세)"
set_run(r, 28, False, RGBColor(0xB7, 0xE3, 0xB0))
t3 = s.shapes.add_textbox(Inches(0.8), Inches(5.5), Inches(11.5), Inches(0.6))
r = t3.text_frame.paragraphs[0].add_run()
r.text = "타임플로어 · 로그뷰어  |  MES / EIF / PLC 로그 분석"
set_run(r, 16, False, RGBColor(0xC5, 0xD4, 0xE8))

# ---- Slide 2: TOC ----
s = new_slide()
add_header_bar(s, "목차")
add_bullets(
    s,
    0.7,
    1.3,
    11.5,
    5.5,
    [
        {"text": "1. 제품 개요 및 구성", "bold": True, "size": 20},
        {"text": "2. 접속 및 실행 방법", "bold": True, "size": 20},
        {"text": "3. 공통 기능 (폴더 선택 / 새로고침)", "bold": True, "size": 20},
        {"text": "4. 타임플로어 상세 사용법", "bold": True, "size": 20},
        {"text": "5. 로그뷰어 상세 사용법", "bold": True, "size": 20},
        {"text": "6. 권장 업무 시나리오", "bold": True, "size": 20},
        {"text": "7. 주의사항 · FAQ", "bold": True, "size": 20},
    ],
)

# ---- Slide 3: Overview ----
s = new_slide()
add_header_bar(s, "1. 제품 개요")
add_bullets(
    s,
    0.6,
    1.2,
    12,
    1.5,
    [
        {
            "text": "EIF Log Tools는 MES · EIF · PLC 관련 로그를 브라우저에서 확인하는 웹 도구입니다.",
            "size": 17,
        },
        {
            "text": "로그는 서버로 업로드되지 않으며, 사용자 PC 브라우저에서만 읽습니다. (Chrome / Edge 권장)",
            "size": 17,
        },
    ],
)
add_card(
    s,
    0.6,
    2.9,
    5.8,
    3.5,
    "타임플로어",
    [
        "· LOT 단위 시퀀스 다이어그램",
        "· MES ↔ EIF ↔ PLC 메시지 흐름",
        "· Alarm 강조 (RESULT / LINESTOP)",
        "· Bit–Word 매칭 상세 표시",
        "· SFC / SOLACE / TRACE 로그 지원",
    ],
)
add_card(
    s,
    6.8,
    2.9,
    5.8,
    3.5,
    "로그뷰어",
    [
        "· 일반 .log 텍스트 뷰어",
        "· 파일/내용 키워드 검색·정렬",
        "· 메시지 블록 단위 표시",
        "· 시간(시:분 ~ 시:분) 필터",
        "· 대용량 로그 빠른 표시",
    ],
)

# ---- Slide 4: Access ----
s = new_slide()
add_header_bar(s, "2. 접속 및 실행 방법")
add_card(
    s,
    0.6,
    1.3,
    5.8,
    5.2,
    "A. 배포 사이트 (Vercel)",
    [
        "1. 제공받은 URL로 접속",
        "   → https://eif-ai-mu.vercel.app",
        "2. 상단 탭에서 기능 선택",
        "   · 타임플로어",
        "   · 로그뷰어",
        "3. 로그 폴더를 선택하고 사용",
        "",
        "※ 인터넷만 되면 설치 없이 사용",
        "※ 로그 파일은 PC에서만 읽힘",
    ],
)
add_card(
    s,
    6.8,
    1.3,
    5.8,
    5.2,
    "B. 로컬 실행 (개발/검증)",
    [
        "1. web 폴더로 이동",
        "2. npm install",
        "3. npm run dev",
        "4. 브라우저에서",
        "   http://localhost:3000 접속",
        "",
        "※ 기능 검증·수정 시 사용",
        "※ 배포 전 로컬에서 먼저 확인",
    ],
)

# ---- Slide 5: Common ----
s = new_slide()
add_header_bar(s, "3. 공통 기능 — 폴더 선택")
add_bullets(
    s,
    0.6,
    1.2,
    12,
    5.5,
    [
        {"text": "로그 폴더 선택 방법", "bold": True, "size": 18},
        {"text": "1. [폴더 선택] 또는 Browse 버튼 클릭", "size": 16},
        {"text": "2. SFC / TRACE / SOLACE 로그가 있는 상위 폴더 지정", "size": 16},
        {"text": "3. 브라우저가 폴더 접근 권한을 요청하면 허용", "size": 16},
        {"text": "", "size": 10},
        {"text": "새로고침", "bold": True, "size": 18},
        {"text": "· [새로고침]으로 같은 폴더의 최신 파일을 다시 읽습니다.", "size": 16},
        {"text": "· Windows 탐색기를 다시 열 필요 없습니다.", "size": 16},
        {"text": "· showDirectoryPicker 지원 브라우저(Chrome/Edge)에서 권한 유지됩니다.", "size": 16},
        {"text": "", "size": 10},
        {"text": "권장 폴더 구조 예: LOG / SFCTYPE / SFC · TRACE", "size": 15, "color": MUTED},
    ],
)

# ---- Slide 6: TimeFloor overview ----
s = new_slide()
add_header_bar(s, "4. 타임플로어 — 화면 구성")
add_bullets(
    s,
    0.6,
    1.2,
    12,
    5.5,
    [
        {"text": "상단: 탭(타임플로어/로그뷰어) · 폴더 선택 · 새로고침", "size": 17},
        {"text": "필터 영역: LOT ID 입력/제안 · From(시작 시각) · 옵션", "size": 17},
        {"text": "본체: MES · EIF · PLC 3레인 시퀀스 다이어그램", "size": 17},
        {"text": "우측/하단: 선택한 이벤트 상세(Word/Bit · Raw)", "size": 17},
        {"text": "", "size": 10},
        {"text": "사용 흐름", "bold": True, "size": 18},
        {"text": "1) 폴더 선택 → 2) LOT 입력/선택 → 3) 타임플로어 생성 → 4) 이벤트 클릭 상세 확인", "size": 16},
    ],
)

# ---- Slide 7: TimeFloor LOT ----
s = new_slide()
add_header_bar(s, "4-1. 타임플로어 — LOT · From 필터")
add_card(
    s,
    0.6,
    1.3,
    6.0,
    5.2,
    "LOT ID",
    [
        "· 입력창에 LOT를 직접 입력하거나",
        "  제안 목록에서 선택합니다.",
        "· 폴더 선택 후 로그에서 LOT를",
        "  추출해 제안합니다.",
        "· LOT 변경 시 이전 결과는 초기화됩니다.",
        "· LOTID가 없는 SFC 메시지도",
        "  주변 맥락에 포함될 수 있습니다.",
    ],
)
add_card(
    s,
    6.9,
    1.3,
    5.7,
    5.2,
    "From (시작 시각)",
    [
        "· 타임플로어에 표시할 시작 시각",
        "· 대문자 From 입력 필드 사용",
        "· 해당 시각 이후 이벤트만 표시",
        "· 긴 로그에서 관심 구간만",
        "  좁힐 때 유용합니다.",
    ],
)

# ---- Slide 8: TimeFloor read ----
s = new_slide()
add_header_bar(s, "4-2. 타임플로어 — 다이어그램 읽기")
add_bullets(
    s,
    0.6,
    1.2,
    12,
    5.5,
    [
        {"text": "레인", "bold": True, "size": 18},
        {"text": "· MES: 상위 시스템 메시지 (SFC 등)", "size": 16},
        {"text": "· EIF: 중계/변환 구간", "size": 16},
        {"text": "· PLC: TRACE Bit/Data 신호", "size": 16},
        {"text": "", "size": 8},
        {"text": "이벤트 확인", "bold": True, "size": 18},
        {"text": "· 화살표/박스를 클릭하면 상세 패널이 열립니다.", "size": 16},
        {"text": "· Bit와 Word가 매칭되면 함께 표시됩니다.", "size": 16},
        {"text": "· Alarm(RESULT / LINESTOP 등)은 강조되어 보입니다.", "size": 16},
        {"text": "· 좌우 스크롤은 다이어그램 영역에만 적용됩니다.", "size": 16},
    ],
)

# ---- Slide 9: TimeFloor tips ----
s = new_slide()
add_header_bar(s, "4-3. 타임플로어 — 실무 팁")
add_bullets(
    s,
    0.6,
    1.2,
    12,
    5.5,
    [
        {"text": "· LOT 제안이 비면: 폴더를 다시 선택하거나 새로고침하세요.", "size": 16},
        {"text": "· Word 상세가 안 보이면: 해당 구간에 Word 로그가 있는지 확인하세요.", "size": 16},
        {"text": "· 결과가 너무 많으면: From 시각을 좁혀 구간을 줄이세요.", "size": 16},
        {"text": "· 다른 LOT로 바꿀 때: LOT를 바꾸면 이전 타임플로어는 자동 초기화됩니다.", "size": 16},
        {"text": "· SFC만 있거나 TRACE만 있어도 가능한 범위에서 표시됩니다.", "size": 16},
    ],
)

# ---- Slide 10: Log Viewer overview ----
s = new_slide()
add_header_bar(s, "5. 로그뷰어 — 화면 구성")
add_bullets(
    s,
    0.6,
    1.2,
    12,
    5.5,
    [
        {"text": "좌측: 로그 파일 목록 (이름 · 검색 · 정렬)", "size": 17},
        {"text": "우측: 선택한 파일 내용 (메시지 블록 단위)", "size": 17},
        {"text": "상단 필터: 키워드 검색 · 시간(HH:mm ~ HH:mm)", "size": 17},
        {"text": "", "size": 10},
        {"text": "사용 흐름", "bold": True, "size": 18},
        {"text": "1) 폴더 선택 → 2) 파일 선택 → 3) 검색/시간 필터 → 4) 블록 단위로 내용 확인", "size": 16},
        {"text": "", "size": 10},
        {"text": "※ 로그뷰어는 로컬/브라우저에서 파일을 직접 읽어 표시합니다.", "size": 15, "color": MUTED},
    ],
)

# ---- Slide 11: Log Viewer search ----
s = new_slide()
add_header_bar(s, "5-1. 로그뷰어 — 검색 · 정렬")
add_card(
    s,
    0.6,
    1.3,
    6.0,
    5.2,
    "검색",
    [
        "· 파일명 / 내용 키워드로 좁힙니다.",
        "· 검색 시 매칭된 SFC/TRACE",
        "  메시지 블록 전체가 펼쳐집니다.",
        "· WORD 등 관련 구간도 함께",
        "  확인할 수 있습니다.",
        "· 대용량 파일은 원문 통째 렌더로",
        "  빠르게 표시합니다.",
    ],
)
add_card(
    s,
    6.9,
    1.3,
    5.7,
    5.2,
    "정렬 · 선택",
    [
        "· 파일 목록을 이름/시간 등으로",
        "  정렬할 수 있습니다.",
        "· 파일 클릭으로 바로 내용을",
        "  불러옵니다.",
        "· Browse가 안 되면 Chrome/Edge로",
        "  다시 시도하세요.",
    ],
)

# ---- Slide 12: Log Viewer time ----
s = new_slide()
add_header_bar(s, "5-2. 로그뷰어 — 시간 필터")
add_bullets(
    s,
    0.6,
    1.2,
    12,
    5.5,
    [
        {"text": "형식: HH:mm ~ HH:mm  (예: 09:30 ~ 10:15)", "bold": True, "size": 18},
        {"text": "· 날짜(하루) 필터는 사용하지 않습니다. 시간대만 지정합니다.", "size": 16},
        {"text": "· 지정한 시·분 구간에 해당하는 메시지 블록만 표시합니다.", "size": 16},
        {"text": "· 검색과 함께 쓰면 ‘키워드 + 시간대’로 빠르게 좁힐 수 있습니다.", "size": 16},
        {"text": "", "size": 10},
        {"text": "메시지 블록", "bold": True, "size": 18},
        {"text": "· 한 건의 SFC/TRACE 메시지를 헤더~본문까지 묶은 단위입니다.",
         "size": 16},
        {"text": "· 검색/시간 필터는 줄 단위가 아니라 블록 단위로 동작합니다.", "size": 16},
    ],
)

# ---- Slide 13: Scenario ----
s = new_slide()
add_header_bar(s, "6. 권장 업무 시나리오")
add_card(
    s,
    0.5,
    1.25,
    4.0,
    5.3,
    "시나리오 A",
    [
        "LOT 트러블 추적",
        "",
        "1. 타임플로어 탭",
        "2. 폴더 선택",
        "3. LOT 입력",
        "4. From으로 구간 축소",
        "5. Alarm/이상 구간 클릭",
        "6. Bit–Word 상세 확인",
    ],
)
add_card(
    s,
    4.7,
    1.25,
    4.0,
    5.3,
    "시나리오 B",
    [
        "원문 로그 확인",
        "",
        "1. 로그뷰어 탭",
        "2. 폴더 선택",
        "3. 파일 검색·선택",
        "4. 키워드 검색",
        "5. 시간 필터 적용",
        "6. 블록 전체 내용 확인",
    ],
)
add_card(
    s,
    8.9,
    1.25,
    3.9,
    5.3,
    "시나리오 C",
    [
        "교차 검증",
        "",
        "1. 타임플로어로",
        "   흐름 파악",
        "2. 시각·메시지명 메모",
        "3. 로그뷰어에서",
        "   같은 시각 검색",
        "4. Raw와 시퀀스 대조",
    ],
)

# ---- Slide 14: Notes ----
s = new_slide()
add_header_bar(s, "7. 주의사항")
add_bullets(
    s,
    0.6,
    1.2,
    12,
    5.5,
    [
        {"text": "· Chrome 또는 Edge를 권장합니다. (폴더 선택 API)", "size": 16},
        {"text": "· 로그는 서버에 업로드되지 않습니다. PC 브라우저 메모리에서만 처리됩니다.", "size": 16},
        {"text": "· 매우 큰 로그는 읽기/표시에 시간이 걸릴 수 있습니다.", "size": 16},
        {"text": "· 폴더 권한을 거부하면 파일을 읽을 수 없습니다. 다시 선택하세요.", "size": 16},
        {"text": "· Vercel 배포본과 로컬(npm run dev) 기능은 동일 계열입니다.", "size": 16},
        {"text": "· 기존 .NET(EIF_AI) 서버 버전과 웹(Next.js) 버전은 별도 실행 경로입니다.", "size": 16},
    ],
)

# ---- Slide 15: FAQ ----
s = new_slide()
add_header_bar(s, "7-1. FAQ")
add_bullets(
    s,
    0.6,
    1.15,
    12,
    5.6,
    [
        {"text": "Q. Browse / 버튼이 안 눌려요", "bold": True, "size": 16},
        {"text": "A. Chrome/Edge로 열고 페이지를 새로고침한 뒤 다시 시도하세요.", "size": 15},
        {"text": "Q. LOT 제안이 안 나와요", "bold": True, "size": 16},
        {"text": "A. 폴더를 다시 선택하거나 [새로고침] 후 LOT 입력창을 확인하세요.", "size": 15},
        {"text": "Q. Word 상세가 비어 있어요", "bold": True, "size": 16},
        {"text": "A. 해당 구간에 Word 로그가 없거나 Bit와 매칭되지 않은 경우입니다.", "size": 15},
        {"text": "Q. 날짜(하루) 필터는 없나요?", "bold": True, "size": 16},
        {"text": "A. 로그뷰어는 HH:mm~HH:mm 시간 필터만 제공합니다.", "size": 15},
        {"text": "Q. 로그가 밖으로 나가나요?", "bold": True, "size": 16},
        {"text": "A. 아니요. 브라우저에서만 읽고 서버로 전송하지 않습니다.", "size": 15},
    ],
)

# ---- Slide 16: End ----
s = new_slide()
bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = NAVY
bg.line.fill.background()
t = s.shapes.add_textbox(Inches(0.8), Inches(2.8), Inches(11.5), Inches(1))
r = t.text_frame.paragraphs[0].add_run()
r.text = "문의 · 개선 요청"
set_run(r, 36, True, WHITE)
t2 = s.shapes.add_textbox(Inches(0.8), Inches(4.0), Inches(11.5), Inches(1.2))
r = t2.text_frame.paragraphs[0].add_run()
r.text = "기능 추가·오류 제보가 있으면 개발 담당자에게 전달해 주세요.\nEIF Log Tools — TimeFloor · Log Viewer"
set_run(r, 18, False, RGBColor(0xC5, 0xD4, 0xE8))

TOTAL = len(prs.slides)
# Cover and end slides: skip numbered footer; middle slides get footer
for idx, slide in enumerate(prs.slides):
    if idx == 0 or idx == TOTAL - 1:
        continue
    add_footer(slide, idx, TOTAL - 1)

prs.save(str(OUT))
print(OUT)
print("slides:", TOTAL)
